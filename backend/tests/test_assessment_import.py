"""Parser/matcher tests. No database: everything here is a pure function over bytes."""
from dataclasses import dataclass
from typing import Optional

import pytest

from app.services.assessment_import import (
    MANAGER_FIELDS,
    SELF_FIELDS,
    ImportError_,
    match_columns,
    normalize_header,
    parse_bytes,
    parse_level,
    resolve_rows,
)

# The engineer-facing self-assessment workbook, header for header. Columns D, E and I
# are printed guidance; only G and H are the engineer's own cells.
SELF_ASSESSMENT_HEADERS = [
    "Skill ID",
    "Domain",
    "Skill",
    "What doing this actually looks like",
    "Evidence that would show it",
    "Team target",
    "MY RATING (0-5)",
    "MY EVIDENCE \u2014 what would you point at?",
    "What my rating means",
]


def self_assessment_bytes(name, rows):
    """A minimal stand-in for Skills_Self_Assessment_INTAKE.xlsx.

    `rows` are (skill_code, guidance, team_target, my_rating, my_evidence) tuples.
    """
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    start = wb.active
    start.title = "Start Here"
    start["A4"], start["B4"] = "Your name", name

    ws = wb.create_sheet("My Assessment")
    ws["A1"] = "My Assessment"
    ws.append([])  # subtitle
    ws.append([])  # spacer
    ws.append(SELF_ASSESSMENT_HEADERS)
    for code, guidance, target, rating, evidence in rows:
        ws.append([code, "AAP", "A skill", "Looks like...", guidance, target, rating, evidence, ""])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- stand-ins for the ORM rows, which is all resolve_rows actually touches ---

@dataclass
class FakeEmployee:
    id: str
    name: str
    email: Optional[str] = None
    employee_id: Optional[str] = None


@dataclass
class FakeSkill:
    id: str
    code: int
    name: str
    target_level: int


@dataclass
class FakeAssessment:
    employee_id: str
    skill_id: str
    self_rating: Optional[int] = None
    reviewer_rating: Optional[int] = None
    target_override: Optional[int] = None
    evidence: Optional[str] = None
    evidence_url: Optional[str] = None


EMPLOYEES = [
    FakeEmployee("emp-001", "Sarah Chen", "sarah.chen@company.com", "ASE-1042"),
    FakeEmployee("emp-002", "Marcus Johnson", "marcus.johnson@company.com", "ASE-1089"),
    # Two people sharing a name — the reason name matching must be able to say "ambiguous".
    FakeEmployee("emp-003", "Alex Twin", "alex.one@company.com", "ASE-1"),
    FakeEmployee("emp-004", "Alex Twin", "alex.two@company.com", "ASE-2"),
]
CATALOG = [
    FakeSkill("aap-01", 1, "AAP architecture", target_level=4),
    FakeSkill("aap-03", 3, "Automation mesh engineering", target_level=4),
]


def csv_bytes(text: str) -> bytes:
    return text.strip().encode()


def resolve(text: str, existing=None, **kwargs):
    rows, _warnings, _source = parse_bytes("t.csv", csv_bytes(text))
    return resolve_rows(rows, EMPLOYEES, CATALOG, existing or [], **kwargs)


# --- header normalization -------------------------------------------------

@pytest.mark.parametrize(
    "raw,expected",
    [
        ("MY EVIDENCE — what would you point at?", "my evidence"),
        ("Key (calculated)", "key"),
        ("  Self   Rating  ", "self rating"),
        ("Evidence / link", "evidence link"),
        ("Critical?", "critical"),
        ("Employee_ID", "employee id"),
    ],
)
def test_normalize_header(raw, expected):
    assert normalize_header(raw) == expected


def test_employee_id_is_not_swallowed_by_employee():
    cols = match_columns(["Employee", "Employee ID", "Skill ID", "Skill"])
    assert cols["employee_name"] == 0
    assert cols["employee_code"] == 1
    assert cols["skill_code"] == 2
    assert cols["skill_name"] == 3


def test_bare_manager_column_is_not_a_rating():
    """The app's own export writes the manager's *name* under a `Manager` header."""
    cols = match_columns(["Employee", "Manager", "Skill ID", "Reviewer rating"])
    assert cols["reviewer_rating"] == 3
    assert cols.get("employee_code") != 1


# --- rating parsing -------------------------------------------------------

@pytest.mark.parametrize("value,expected", [(0, 0), (5, 5), ("3", 3), ("Independent", 3), ("expert", 4)])
def test_parse_level_accepts(value, expected):
    level, is_clear, error = parse_level(value)
    assert (level, is_clear, error) == (expected, False, None)


@pytest.mark.parametrize("value", [6, 7, -1, "45", "banana"])
def test_parse_level_rejects_instead_of_clamping(value):
    """lib/utils.ts clampLevel() turns a typo'd 45 into a 5. An import must not."""
    level, _is_clear, error = parse_level(value)
    assert level is None
    assert error


def test_parse_level_blank_is_unchanged_and_dash_is_clear():
    assert parse_level("") == (None, False, None)
    assert parse_level(None) == (None, False, None)
    assert parse_level("-")[1] is True
    assert parse_level("clear")[1] is True


# --- row statuses ---------------------------------------------------------

def test_matches_by_email_employee_id_and_name():
    result = resolve(
        """
Email,Employee ID,Employee,Skill ID,Self rating
sarah.chen@company.com,,,1,3
,ASE-1089,,1,3
,,Sarah Chen,3,3
"""
    )
    assert [r.matched_by for r in result.rows] == ["email", "employee_id", "name"]
    assert [r.employee_id for r in result.rows] == ["emp-001", "emp-002", "emp-001"]


def test_duplicate_last_row_wins():
    result = resolve(
        """
Email,Skill ID,Self rating
sarah.chen@company.com,1,2
sarah.chen@company.com,1,4
"""
    )
    assert [r.status for r in result.rows] == ["duplicate", "ok"]
    assert result.rows[1].values["self_rating"] == 4


def test_ambiguous_name_is_not_a_match():
    result = resolve("Employee,Skill ID,Self rating\nAlex Twin,1,3")
    assert result.rows[0].status == "ambiguous_employee"
    assert result.rows[0].employee_id is None


def test_unknown_employee_and_skill():
    result = resolve(
        """
Email,Skill ID,Self rating
nobody@company.com,1,3
sarah.chen@company.com,9999,3
"""
    )
    assert [r.status for r in result.rows] == ["unknown_employee", "unknown_skill"]


def test_invalid_row_still_reports_who_it_was():
    result = resolve("Email,Skill ID,Self rating\nsarah.chen@company.com,1,7")
    row = result.rows[0]
    assert row.status == "invalid_value"
    assert row.employee_name == "sarah.chen@company.com"
    assert "0-5" in row.messages[0]


def test_unchanged_when_file_agrees_with_stored():
    existing = [FakeAssessment("emp-001", "aap-01", self_rating=3)]
    result = resolve("Email,Skill ID,Self rating\nsarah.chen@company.com,1,3", existing)
    assert result.rows[0].status == "unchanged"


def test_blank_cell_leaves_a_field_alone():
    result = resolve("Email,Skill ID,Self rating,Reviewer rating\nsarah.chen@company.com,1,3,")
    assert result.rows[0].values == {"self_rating": 3}


def test_dash_clears_a_stored_value():
    existing = [FakeAssessment("emp-001", "aap-01", self_rating=3)]
    result = resolve("Email,Skill ID,Self rating\nsarah.chen@company.com,1,-", existing)
    assert result.rows[0].values == {"self_rating": None}
    assert result.rows[0].status == "ok"


# --- the target_override correctness rule ---------------------------------

def test_target_matching_the_catalog_does_not_write_an_override():
    """effectiveTarget() is `targetOverride ?? def.targetLevel`.

    The export writes a Target on every row, so storing it verbatim would pin an
    override on every assessment and permanently detach people from catalog targets.
    """
    result = resolve("Email,Skill ID,Target,Self rating\nsarah.chen@company.com,1,4,3")
    assert result.rows[0].values["target_override"] is None


def test_target_differing_from_the_catalog_is_kept():
    result = resolve("Email,Skill ID,Target,Self rating\nsarah.chen@company.com,1,2,3")
    assert result.rows[0].values["target_override"] == 2


# --- field-level authorization --------------------------------------------

def test_employee_may_not_set_reviewer_rating():
    result = resolve(
        "Email,Skill ID,Self rating,Reviewer rating\nsarah.chen@company.com,1,3,5",
        allowed_fields=SELF_FIELDS,
    )
    row = result.rows[0]
    assert "reviewer_rating" not in row.values
    assert row.values["self_rating"] == 3
    assert any("not yours to set" in m for m in row.messages)


def test_manager_may_set_reviewer_rating():
    result = resolve(
        "Email,Skill ID,Self rating,Reviewer rating\nsarah.chen@company.com,1,3,5",
        allowed_fields=MANAGER_FIELDS,
    )
    assert result.rows[0].values["reviewer_rating"] == 5


# --- file-level errors ----------------------------------------------------

def test_missing_employee_column_is_rejected():
    with pytest.raises(ImportError_, match="No employee column"):
        parse_bytes("t.csv", csv_bytes("Skill ID,Self rating\n1,3"))


def test_unsupported_extension_is_rejected():
    with pytest.raises(ImportError_) as exc:
        parse_bytes("legacy.xls", b"anything")
    assert exc.value.status_code == 415


def test_tab_separated_file_is_read():
    rows, _w, source = parse_bytes(
        "t.tsv", b"Email\tSkill ID\tSelf rating\nsarah.chen@company.com\t1\t3"
    )
    assert source == "csv"
    assert rows[0].values["self_rating"] == 3


# --- guidance columns -----------------------------------------------------

def test_guidance_columns_are_never_bound():
    """`Evidence that would show it` is rubric text, printed the same for everyone.

    It sits three columns left of the engineer's own `MY EVIDENCE`, so the loose
    `evidence` keyword fallback used to claim it and `_row_from_cells` then let it
    supersede the real answer — giving every team member identical evidence.
    """
    cols = match_columns(SELF_ASSESSMENT_HEADERS)
    assert cols["self_evidence"] == 7
    assert cols.get("evidence") is None
    assert cols.get("skill_name") == 2
    assert cols.get("self_rating") == 6


def test_catalog_reference_columns_are_not_bound():
    """The Skill Catalog export describes each skill; none of it is an input cell."""
    cols = match_columns(
        ["Skill ID", "Skill", "Observable capability", "Example evidence", "Target level"]
    )
    assert cols.get("evidence") is None
    assert cols["skill_code"] == 0
    assert cols["target"] == 4


def test_named_evidence_column_still_wins_over_self_evidence():
    """The manager sheet's calibrated evidence must keep superseding the self answer."""
    cols = match_columns(["Employee", "Skill ID", "My evidence", "Reviewer evidence"])
    assert cols["self_evidence"] == 2
    assert cols["evidence"] == 3


def test_self_assessment_keeps_each_engineers_own_evidence():
    data = self_assessment_bytes(
        "Sarah Chen",
        [
            (1, "Architecture decision and topology review", 4, 3, "Rebuilt the mesh topology"),
            (3, "Mesh design and failure diagnosis", 4, 2, None),
        ],
    )
    rows, warnings, source = parse_bytes("intake.xlsx", data)
    assert source == "xlsx"
    assert warnings == []
    assert [r.employee_name for r in rows] == ["Sarah Chen", "Sarah Chen"]
    assert rows[0].values["evidence"] == "Rebuilt the mesh topology"
    # Blank means "leave unchanged" — not "overwrite with the guidance text".
    assert "evidence" not in rows[1].values


def test_self_assessment_does_not_import_the_printed_team_target():
    """`Team target` is the catalog level shown for context, not an override the

    engineer can set. Importing it would silently drop one the manager had set.
    """
    data = self_assessment_bytes("Sarah Chen", [(1, "Guidance", 4, 3, "Evidence")])
    rows, _warnings, _source = parse_bytes("intake.xlsx", data)
    assert "target_override" not in rows[0].values
    assert rows[0].values == {"self_rating": 3, "evidence": "Evidence"}


def test_untouched_self_assessment_rows_are_dropped_not_imported():
    data = self_assessment_bytes("Sarah Chen", [(1, "Guidance", 4, None, None)])
    rows, warnings, _source = parse_bytes("intake.xlsx", data)
    assert rows == []
    assert any("Nothing was filled in" in w for w in warnings)


def test_reimporting_the_same_self_assessment_reports_no_change():
    """Re-uploading to add evidence must not re-write ratings that already agree."""
    data = self_assessment_bytes("Sarah Chen", [(1, "Guidance", 4, 3, "Rebuilt the mesh")])
    rows, _warnings, _source = parse_bytes("intake.xlsx", data)
    stored = [FakeAssessment("emp-001", "aap-01", self_rating=3, evidence="Rebuilt the mesh")]
    result = resolve_rows(rows, EMPLOYEES, CATALOG, stored)
    assert result.rows[0].status == "unchanged"
